"""独立逐项核验账本、版本和模板工作簿。"""
from pathlib import Path
import json
import numpy as np
import pandas as pd
from openpyxl import load_workbook
R=Path(__file__).resolve().parent
d=pd.read_csv(R/'dispatch_all.csv'); x=pd.read_csv(R/'dispatch.csv'); v=pd.read_csv(R/'schedule_versions.csv')
checks={}
def eq(name,a,b,tol=1e-5):
    error=float(np.max(np.abs(np.asarray(a,dtype=float)-np.asarray(b,dtype=float))))
    assert error<tol,(name,error)
    checks[name]=error
assert len(d)==52560 and len(x)==48096 and x.date.nunique()==334 and len(v)==120240
eq('balance',d.adjusted_grid_kwh+d.pv_kwh+d.discharge_kwh+d.emergency_kwh,d.load_kwh+d.charge_kwh+d.unused_kwh)
eq('state',d.storage_end_kwh,d.storage_start_kwh+.9*d.charge_kwh-d.discharge_kwh/.9)
eq('continuity',d.storage_start_kwh.iloc[1:],d.storage_end_kwh.iloc[:-1])
assert d.storage_end_kwh.between(1200-1e-5,10800+1e-5).all()
assert d[['charge_kwh','discharge_kwh']].max().max()<=5000/6+1e-5
assert d[['original_grid_kwh','adjusted_grid_kwh','charge_kwh','discharge_kwh','emergency_kwh','unused_kwh']].min().min()>=-1e-5
eq('mutual_exclusion',np.minimum(d.charge_kwh,d.discharge_kwh),0)
assert (d.charge_kwh<=d.adjusted_grid_kwh+d.pv_kwh+1e-5).all()
eq('increase',d.increase_kwh,np.maximum(d.adjusted_grid_kwh-d.original_grid_kwh,0))
eq('decrease',d.decrease_kwh,np.maximum(d.original_grid_kwh-d.adjusted_grid_kwh,0))
eq('total_cost',d.total_cost_yuan,d.price_yuan_per_kwh*(d.original_grid_kwh+1.5*d.increase_kwh+.5*d.decrease_kwh+5*d.emergency_kwh))
assert (v.forecast_release_hour<=v.decision_hour).all() and (v.slot>v.decision_hour*6).all()
latest=v.sort_values('decision_hour').drop_duplicates(['date','slot'],keep='last').sort_values(['date','slot'])
eq('final_version',latest.version_grid_kwh,x.adjusted_grid_kwh)
zero=v[v.decision_hour==0].sort_values(['date','slot'])
eq('original_version',zero.version_grid_kwh,x.original_grid_kwh)
wb=load_workbook(R/'result3.xlsx',data_only=True)
for name,col,cost in [('计划购电量','original_grid_kwh','plan_cost_yuan'),('调整购电量','adjusted_grid_kwh',None)]:
    ws=wb[name];assert ws.max_row==335
    eq(name,[ws.cell(i,j).value for i in range(2,336) for j in range(2,146)],x[col])
    for i,(date,g) in enumerate(x.groupby('date'),2):
        assert str(ws.cell(i,1).value.date())==date
        eq(name+date+'sum',ws.cell(i,146).value,g[col].sum())
        eq(name+date+'cost',ws.cell(i,147).value,g[cost].sum() if cost else (g.increase_cost_yuan+g.decrease_cost_yuan).sum())
ws=wb['充放电量'];row=2
for date,g in x.groupby('date'):
    for k in range(6):
        z=g.iloc[k*24:(k+1)*24]
        eq('charge'+str(row),ws.cell(row,3).value,z.charge_kwh.sum());eq('discharge'+str(row),ws.cell(row,4).value,z.discharge_kwh.sum())
        if k<2:eq('endpoint'+str(row),ws.cell(row,6).value,g.storage_start_kwh.iloc[0] if k==0 else g.storage_end_kwh.iloc[-1])
        row+=1
events=pd.read_csv(R/'emergency_events.csv');ws=wb['紧急购电量']
eq('emergency_excel',[ws.cell(i,3).value for i in range(2,len(events)+2)],events.emergency_kwh)
eq('emergency_total',events.emergency_kwh.sum(),x.emergency_kwh.sum(),1e-3)
for i,z in enumerate(events.itertuples(),2):
    assert str(ws.cell(i,1).value.date())==z.date and ws.cell(i,2).value==z.interval
comparison=pd.read_csv(R/'policy_comparison.csv')
assert len(comparison)==12
eq('fair_initial',comparison[comparison.fair_boundary].initial_kwh,6000)
eq('fair_final',comparison[comparison.fair_boundary].final_kwh,6000)
report={'status':'PASS','period_rows':len(x),'version_rows':len(v),'checks':len(checks),'max_error':max(checks.values()),'workbook_sheets':wb.sheetnames}
(R/'results/verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(report)

