"""问题三：按已发布预报滚动调整购电，完整保留0时承诺和调整版本。"""
from pathlib import Path
from datetime import datetime
from concurrent.futures import ProcessPoolExecutor
import argparse, json, hashlib, time
import numpy as np
import pandas as pd
import openpyxl
from scipy.optimize import linprog
from scipy.sparse import lil_matrix
from battery_core import (INPUT, DATES, SELECTED, EMIN, EMAX, EREF, POWER, T,
                          read_data, minute, predict, plan_day, value_table, action, interval, episodes)

ROOT=Path(__file__).resolve().parent
POLICIES={
    'only0':((0,),{}),
    'add6':((0,1),{}),
    'add12':((0,1,2),{}),
    'all':((0,1,2,3),{}),
    'stale_all':((0,1,2,3),{1:0,2:0,3:0}),
    'stale18':((0,1,2,3),{3:2}),
}

def read_forecasts(folder,actual):
    w=openpyxl.load_workbook(folder/'附件3.xlsx',read_only=True,data_only=True)
    rows=list(w.active.values);w.close()
    assert len(rows)==1461 and all(len(r)==26 for r in rows)
    assert list(rows[0])==['日期','预报时刻']+[f'预报{i}小时' for i in range(1,25)]
    forecast=np.full((365,4,T),np.nan);hourly=np.empty((365,4,24));date=None
    for k,row in enumerate(rows[1:]):
        d,r=divmod(k,4);offset=r*36
        if row[0] not in ('',None):date=pd.Timestamp(row[0])
        assert date==DATES[d] and minute(row[1])==r*360
        values=np.asarray(row[2:],float)
        assert np.isfinite(values).all() and (values>=0).all()
        hourly[d,r]=values
        # 发布时刻的观测是上一10分钟段的右端点；0时只用昨日末值。
        anchor=actual[d,1,offset-1]*6 if r else (actual[d-1,1,-1]*6 if d else 0.)
        edges=np.interp(np.arange(T+1)/6,np.arange(25),np.r_[anchor,values])
        energy=(edges[:-1]+edges[1:])/12
        forecast[d,r,offset:]=energy[:T-offset]
    return forecast,hourly

def historical_predictions(actual):
    pred=np.zeros((365,T))
    for d in range(1,365):pred[d]=predict(actual[:d],DATES[d],7)[0]
    return pred

def make_scene(actual,forecasts,load_pred,d,source_release,offset,count=7):
    """每个联合残差来自同一个已经结束的历史日；日内以前发布的预报仍可取后缀。"""
    assert d>0 and source_release*36<=offset
    ids=np.arange(max(1,d-count),d)
    base=np.stack([load_pred[d,offset:],forecasts[d,source_release,offset:]])
    if len(ids)==0:return base[None,:,:]
    load=base[0]+actual[ids,0,offset:]-load_pred[ids,offset:]
    pv=base[1]+actual[ids,1,offset:]-forecasts[ids,source_release,offset:]
    scene=np.maximum(np.stack([load,pv],axis=1),0)
    assert np.isfinite(scene).all()
    return scene

def adjust_plan(p,scene,initial,original,eta=.9,kappa=.472,refund=False,hard_terminal=False):
    """以0时原计划为基准结算，当前窗口原计划费用为常数，不重复计入优化目标。"""
    s,n=scene.shape[0],len(p)
    a,c,b,e=[np.arange(k*n,(k+1)*n) for k in range(4)]
    q=np.arange(4*n,(4+s)*n).reshape(s,n)
    u=np.arange((4+s)*n,(5+s)*n);v=np.arange((5+s)*n,(6+s)*n)
    h=(6+s)*n;nv=h+1
    obj=np.zeros(nv);obj[u]=1.5*p;obj[v]=(-.5 if refund else .5)*p
    obj[q]=5*p[None,:]/s;obj[h]=kappa;obj[c]=obj[b]=1e-8
    bounds=[(0,None)]*nv
    for t in range(n):
        bounds[c[t]]=bounds[b[t]]=(0,POWER/1000)
        bounds[e[t]]=(EMIN/1000,EMAX/1000)
    if hard_terminal:
        bounds[e[-1]]=(EREF/1000,EREF/1000)
        for t in range(max(0,n-7),n):bounds[a[t]]=(POWER/1000,None)
    eq=lil_matrix((2*n,nv));rhs=np.zeros(2*n)
    for t in range(n):
        eq[t,e[t]]=1;eq[t,c[t]]=-eta;eq[t,b[t]]=1/eta
        if t:eq[t,e[t-1]]=-1
        else:rhs[t]=initial/1000
        eq[n+t,a[t]]=1;eq[n+t,u[t]]=-1;eq[n+t,v[t]]=1;rhs[n+t]=original[t]/1000
    mat=lil_matrix((s*n+n+1,nv));ub=np.zeros(s*n+n+1)
    for j in range(s):
        for t in range(n):
            row=j*n+t
            mat[row,a[t]]=-1;mat[row,b[t]]=-1;mat[row,c[t]]=1;mat[row,q[j,t]]=-1
            ub[row]=(scene[j,1,t]-scene[j,0,t])/1000
    for t in range(n):
        mat[s*n+t,c[t]]=1;mat[s*n+t,a[t]]=-1;ub[s*n+t]=scene[:,1,t].min()/1000
    mat[-1,e[-1]]=-1;mat[-1,h]=-1;ub[-1]=-EREF/1000
    res=linprog(obj,A_ub=mat.tocsr(),b_ub=ub,A_eq=eq.tocsr(),b_eq=rhs,bounds=bounds,method='highs',
                options={'primal_feasibility_tolerance':1e-8,'dual_feasibility_tolerance':1e-8})
    if not res.success:raise RuntimeError(res.message)
    x=res.x*1000;grid=np.maximum(x[a],0)
    cancel=np.minimum(x[c],x[b]/eta**2)
    charge=np.maximum(x[c]-cancel,0);discharge=np.maximum(x[b]-eta**2*cancel,0)
    assert np.max(np.abs(np.r_[initial,x[e][:-1]]+eta*charge-discharge/eta-x[e]))<1e-4
    assert np.max(np.minimum(charge,discharge))<1e-5
    assert np.max(np.maximum(charge-grid-scene[:,1,:].min(axis=0),0))<1e-4
    if not refund:assert (grid>=original-1e-5).all(), 'Non-refund reduction contradicts free-surplus dominance'
    return grid,x[e],float(res.fun*1000)

def simulate_day(p,actual,forecasts,load_pred,d,initial,policy='all',step=100.,eta=.9,
                 count=7,kappa_multiplier=1.,refund=False,hard_terminal=False,keep_versions=False):
    active,source_override=POLICIES[policy];kappa=p[0]/eta*kappa_multiplier
    energy=initial;results=[];versions=[];grid=None;reference=None
    active_grid=np.zeros(T);original=np.zeros(T);ref_all=np.full(T,initial)
    source_all=np.zeros(T,int);decision_all=np.zeros(T,int)
    objective=[]
    if d==0:
        q=np.maximum(actual[d,0]-actual[d,1],0);w=np.maximum(actual[d,1]-actual[d,0],0)
        out=np.column_stack([np.full(T,initial),np.full(T,initial),np.zeros((T,2)),q,w])
    else:
        for t in range(T):
            r=t//36
            if t%36==0 and r in active:
                source=source_override.get(r,r);scene=make_scene(actual,forecasts,load_pred,d,source,t,count)
                if r==0:
                    grid,reference,info=plan_day(p,scene,energy,eta,kappa,hard_terminal)
                    original=grid.copy();obj=info['lp_objective_yuan']
                else:
                    grid,reference,obj=adjust_plan(p[t:],scene,energy,original[t:],eta,kappa,refund,hard_terminal)
                objective.append(obj);start=t
                states,values=value_table(p[t:],grid,scene,eta,kappa,step,hard_terminal)
                active_grid[t:]=grid;ref_all[t:]=reference;source_all[t:]=source*6;decision_all[t:]=r*6
                if keep_versions:
                    for k in range(T-t):
                        versions.append((str(DATES[d].date()),r*6,source*6,t+k+1,float(original[t+k]),float(grid[k]),
                                         float(reference[k]),float(load_pred[d,t+k]),float(forecasts[d,source,t+k])))
            energy,c,b,q,w=action(p[t],grid[t-start],actual[d,0,t],actual[d,1,t],energy,states,values[t-start+1],eta)
            results.append((initial if not results else results[-1][1],energy,c,b,q,w))
        out=np.asarray(results)
    inc=np.maximum(active_grid-original,0);dec=np.maximum(original-active_grid,0)
    df=pd.DataFrame(out,columns=['storage_start_kwh','storage_end_kwh','charge_kwh','discharge_kwh','emergency_kwh','unused_kwh'])
    df.insert(0,'date',str(DATES[d].date()));df.insert(1,'slot',np.arange(1,T+1))
    df['interval']=[interval(t) for t in range(T)]
    df['price_yuan_per_kwh']=p;df['load_kwh']=actual[d,0];df['pv_kwh']=actual[d,1]
    df['original_grid_kwh']=original;df['adjusted_grid_kwh']=active_grid
    df['increase_kwh']=inc;df['decrease_kwh']=dec;df['reference_storage_kwh']=ref_all
    df['forecast_release_hour']=source_all;df['decision_hour']=decision_all
    df['plan_cost_yuan']=p*original;df['increase_cost_yuan']=1.5*p*inc
    df['decrease_cost_yuan']=(-.5 if refund else .5)*p*dec
    df['emergency_cost_yuan']=5*p*df.emergency_kwh
    df['total_cost_yuan']=df[['plan_cost_yuan','increase_cost_yuan','decrease_cost_yuan','emergency_cost_yuan']].sum(axis=1)
    df['forecast_load_kwh']=load_pred[d]
    df['forecast_pv_kwh']=[forecasts[d,int(source_all[t]/6),t] for t in range(T)]
    df['history_latest_date']='' if d==0 else str(DATES[d-1].date())
    daily={'date':str(DATES[d].date()),'policy':policy,'initial_kwh':initial,'final_kwh':float(out[-1,1]),
           'emergency_slots':int((df.emergency_kwh>1e-6).sum()),'solve_count':len(objective)}
    for col in ['original_grid_kwh','adjusted_grid_kwh','increase_kwh','decrease_kwh','emergency_kwh','charge_kwh',
                'discharge_kwh','unused_kwh','plan_cost_yuan','increase_cost_yuan','decrease_cost_yuan','emergency_cost_yuan','total_cost_yuan']:
        daily[col]=float(df[col].sum())
    return df,daily,versions

def validate(df,eta=.9,refund=False):
    err={
        'balance_kwh':float(abs(df.adjusted_grid_kwh+df.pv_kwh+df.discharge_kwh+df.emergency_kwh-df.load_kwh-df.charge_kwh-df.unused_kwh).max()),
        'dynamics_kwh':float(abs(df.storage_start_kwh+eta*df.charge_kwh-df.discharge_kwh/eta-df.storage_end_kwh).max()),
        'continuity_kwh':float(abs(df.storage_start_kwh.to_numpy()[1:]-df.storage_end_kwh.to_numpy()[:-1]).max()),
        'simultaneous_kwh':float(np.minimum(df.charge_kwh,df.discharge_kwh).max()),
        'charge_source_kwh':float(np.maximum(df.charge_kwh-df.adjusted_grid_kwh-df.pv_kwh,0).max()),
        'power_kw':float(max(0,df[['charge_kwh','discharge_kwh']].to_numpy().max()*6-5000)),
    }
    assert max(err.values())<1e-4,err
    assert df[['storage_start_kwh','storage_end_kwh']].min().min()>=EMIN-1e-5
    assert df[['storage_start_kwh','storage_end_kwh']].max().max()<=EMAX+1e-5
    assert np.isfinite(df.select_dtypes('number')).all().all()
    assert (df.forecast_release_hour<=df.decision_hour).all()
    assert (df.decision_hour*6<=df.slot-1).all()
    assert abs(df.total_cost_yuan.sum()-df[['plan_cost_yuan','increase_cost_yuan','decrease_cost_yuan','emergency_cost_yuan']].to_numpy().sum())<1e-5
    if not refund:assert df.decrease_kwh.max()<1e-5
    return err

def run_case(payload):
    p,actual,forecasts,load_pred,days,policy,fair=payload
    energy=EREF;tables=[];daily=[];versions=[]
    start=31 if fair else 0
    for d in range(start,days):
        table,ds,vs=simulate_day(p,actual,forecasts,load_pred,d,energy,policy,
                               hard_terminal=fair and d==days-1,keep_versions=policy=='all' and not fair)
        energy=ds['final_kwh'];tables.append(table);daily.append(ds);versions.extend(vs)
    table=pd.concat(tables,ignore_index=True);checks=validate(table)
    if fair:assert abs(energy-EREF)<1e-5
    all_daily=pd.DataFrame(daily)
    # 非主策略只保存每天汇总，物理约束仍检查全部10分钟执行轨迹。
    return {'policy':policy,'fair':fair,'daily':all_daily,'table':table if policy=='all' and not fair else None,
            'versions':versions,'validation':checks}

def write_excel(folder,out,df,daily,events):
    wb=openpyxl.load_workbook(folder/'附件5/result3.xlsx')
    for name,col,totalcol,costcol in [('计划购电量','original_grid_kwh','original_grid_kwh','plan_cost_yuan'),
                                       ('调整购电量','adjusted_grid_kwh','adjusted_grid_kwh',None)]:
        ws=wb[name]
        for t in range(T):ws.cell(1,t+2,interval(t))
        ws.cell(1,146,'全天电量（kWh）');ws.cell(1,147,'原计划费用（元）' if costcol else '调整相关费用（元）')
        for row,(date,day) in enumerate(df.groupby('date',sort=False),2):
            ws.cell(row,1,datetime.fromisoformat(date))
            for colno,value in enumerate(day[col],2):ws.cell(row,colno,float(value))
            ws.cell(row,146,float(day[totalcol].sum()))
            ws.cell(row,147,float(day[costcol].sum() if costcol else day.increase_cost_yuan.sum()+day.decrease_cost_yuan.sum()))
        ws.freeze_panes='B2'
    ws=wb['充放电量']
    for m in list(ws.merged_cells.ranges):ws.unmerge_cells(str(m))
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.value=None
    row=2
    for date,day in df.groupby('date',sort=False):
        for k in range(6):
            part=day.iloc[k*24:(k+1)*24]
            values=[datetime.fromisoformat(date),f'{4*k:02d}:00-{4*k+4:02d}:00',float(part.charge_kwh.sum()),float(part.discharge_kwh.sum())]
            if k==0:values+=['00:00',float(day.storage_start_kwh.iloc[0])]
            if k==1:values+=['24:00',float(day.storage_end_kwh.iloc[-1])]
            for c,value in enumerate(values,1):ws.cell(row,c,value)
            row+=1
    ws=wb['紧急购电量']
    for m in list(ws.merged_cells.ranges):ws.unmerge_cells(str(m))
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.value=None
    for row,event in enumerate(events.itertuples(),2):
        ws.cell(row,1,datetime.fromisoformat(event.date));ws.cell(row,2,event.interval);ws.cell(row,3,float(event.emergency_kwh))
    ws=wb.create_sheet('每日费用汇总');ws.append(list(daily.columns))
    for row in daily.itertuples(index=False,name=None):ws.append(row)
    ws=wb.create_sheet('口径说明')
    for row in [('项目','说明'),('时间','前一10分钟均值解释；输出覆盖00:00至24:00，原模板标签在副本修正。'),
                ('计划与调整','计划表为0时承诺G；调整表为每段执行前最终有效普通购电总量A，不是增购量。'),
                ('主计费','原计划全额计费；总费用=pG+1.5p(A-G)+ +0.5p(G-A)+ +5pe。原计划不退款。'),
                ('多次调整','每段最终A与原0时G比较一次收费；中间版本不累加收费。'),
                ('减购','允许免费剩余时，主口径最优解不会A<G；后次版本可撤回之前额外增加的部分。'),
                ('应急','应急电量另列，调整表不含应急；当前段可测供需并即时响应，未模拟测量延迟。'),
                ('储能','实际执行量；每日期初等于前日期末；只有1月1日初始6000。两侧效率分别0.9。'),
                ('信息','只使用已发布光伏预报和以前完整日残差；没有把未来实际值用于日前计划。'),
                ('算法','共享场景LP与100kWh网格价值反馈，属于顺序近似策略，无一般多阶段全局最优声明。'),
                ('版本','全部0/6/12/18时预报及购电版本另存schedule_versions.csv。')]:ws.append(row)
    for ws in wb:
        ws.column_dimensions['A'].width=20
        for row in ws:
            for cell in row:
                if isinstance(cell.value,datetime):cell.number_format='yyyy-mm-dd'
                elif isinstance(cell.value,(float,int)):cell.number_format='0.000000'
    for name in ['充放电量','紧急购电量']:
        for c in ['B','C','D','E','F']:wb[name].column_dimensions[c].width=21
    wb['口径说明'].column_dimensions['B'].width=110
    wb.save(out/'result3.xlsx')

def forecast_diagnostics(actual,forecasts):
    rows=[]
    # 所有版本只比较共同的13:00—18:00区间，避免提前量和夜间零值混淆。
    for r in [0,1,2]:
        error=forecasts[31:,r,78:108]-actual[31:,1,78:108]
        for d in range(334):
            rows.append({'date':str(DATES[d+31].date()),'release_hour':r*6,'comparison_window':'13:00-18:00',
                         'mae_kw':float(np.abs(error[d]).mean()*6),'rmse_kw':float(np.sqrt(np.mean(error[d]**2))*6)})
    return pd.DataFrame(rows)

def bootstrap_comparisons(fair_daily):
    pairs=[('only0','add6'),('add6','add12'),('add12','all'),('only0','all'),('stale_all','all'),('stale18','all')]
    rng=np.random.default_rng(2026);rows=[]
    for before,after in pairs:
        x=fair_daily[before].total_cost_yuan.to_numpy()-fair_daily[after].total_cost_yuan.to_numpy()
        for block in [3,7,14]:
            starts=rng.integers(0,len(x),size=(2000,int(np.ceil(len(x)/block))))
            indices=(starts[:,:,None]+np.arange(block))%len(x)
            samples=x[indices.reshape(2000,-1)[:,:len(x)]].mean(axis=1)
            lo,hi=np.quantile(samples,[.025,.975])
            rows.append({'before':before,'after':after,'block_days':block,'mean_saving_yuan_per_day':float(x.mean()),
                         'total_saving_yuan':float(x.sum()),'ci_low_yuan_per_day':float(lo),'ci_high_yuan_per_day':float(hi),
                         'bootstrap_repeats':2000,'seed':2026})
    return pd.DataFrame(rows)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--input',type=Path,default=INPUT)
    parser.add_argument('--output',type=Path,default=ROOT);parser.add_argument('--days',type=int,default=365)
    parser.add_argument('--full',action='store_true');parser.add_argument('--workers',type=int,default=2)
    args=parser.parse_args();begin=time.perf_counter();out=args.output.resolve()
    assert 2<=args.days<=365 and 1<=args.workers<=4
    assert out!=args.input.resolve() and args.input.resolve() not in out.parents
    out.mkdir(parents=True,exist_ok=True);(out/'results').mkdir(exist_ok=True)
    p,actual=read_data(args.input);forecast,hourly=read_forecasts(args.input,actual);load_pred=historical_predictions(actual)
    cases=[('all',False)]
    if args.full:
        assert args.days==365
        cases=[(name,fair) for fair in [False,True] for name in POLICIES]
    payload=[(p,actual,forecast,load_pred,args.days,name,fair) for name,fair in cases]
    results=[]
    if args.full:
        with ProcessPoolExecutor(max_workers=args.workers) as pool:
            for res in pool.map(run_case,payload):
                results.append(res);print('Completed',res['policy'],'fair' if res['fair'] else 'main',flush=True)
    else:results=[run_case(payload[0])]
    main_result=next(r for r in results if r['policy']=='all' and not r['fair'])
    all_table=main_result['table'];all_daily=main_result['daily']
    all_table.to_csv(out/'dispatch_all.csv',index=False,encoding='utf-8-sig')
    all_daily.to_csv(out/'daily_all.csv',index=False,encoding='utf-8-sig')
    df=all_table[all_table.date>='2025-02-01'].copy();daily=all_daily[all_daily.date>='2025-02-01'].copy()
    versions=pd.DataFrame(main_result['versions'],columns=['date','decision_hour','forecast_release_hour','slot','original_grid_kwh',
                                                         'version_grid_kwh','reference_storage_kwh','forecast_load_kwh','forecast_pv_kwh'])
    versions.to_csv(out/'schedule_versions_all.csv',index=False,encoding='utf-8-sig')
    summary={'days_run':args.days,'parameters':{'window_days':7,'scenario_count':7,'state_grid_kwh':100,'eta':.9,'kappa':float(p[0]/.9),
             'settlement':'original fully charged; final-version net adjustment charged once','evaluation':'2025-02-01 through 2025-12-31'},
             'input_hashes':{name:hashlib.sha256((args.input/name).read_bytes()).hexdigest() for name in ['附件1.xlsx','附件2.xlsx','附件3.xlsx','附件5/result3.xlsx']},
             'validation':main_result['validation']}
    if len(df):
        events=episodes(df);df.to_csv(out/'dispatch.csv',index=False,encoding='utf-8-sig');daily.to_csv(out/'daily.csv',index=False,encoding='utf-8-sig')
        events.to_csv(out/'emergency_events.csv',index=False,encoding='utf-8-sig')
        versions[versions.date>='2025-02-01'].to_csv(out/'schedule_versions.csv',index=False,encoding='utf-8-sig')
        write_excel(args.input,out,df,daily,events)
        summary['evaluation']={c:float(daily[c].sum()) for c in daily.select_dtypes('number') if c not in ['initial_kwh','final_kwh']}
        summary['evaluation'].update({'days':len(daily),'initial_kwh':float(daily.initial_kwh.iloc[0]),'final_kwh':float(daily.final_kwh.iloc[-1]),'emergency_events':len(events)})
    if args.full:
        rows=[];fair_daily={}
        for res in results:
            ds=res['daily'];evaluation=ds[ds.date>='2025-02-01']
            if res['fair']:fair_daily[res['policy']]=evaluation
            name=f"{'fair' if res['fair'] else 'policy'}_{res['policy']}_daily.csv"
            evaluation.to_csv(out/name,index=False,encoding='utf-8-sig')
            rows.append({'policy':res['policy'],'fair_boundary':res['fair'],'total_cost_yuan':float(evaluation.total_cost_yuan.sum()),
                         'plan_cost_yuan':float(evaluation.plan_cost_yuan.sum()),'adjustment_cost_yuan':float(evaluation.increase_cost_yuan.sum()+evaluation.decrease_cost_yuan.sum()),
                         'emergency_cost_yuan':float(evaluation.emergency_cost_yuan.sum()),'emergency_kwh':float(evaluation.emergency_kwh.sum()),
                         'initial_kwh':float(evaluation.initial_kwh.iloc[0]),'final_kwh':float(evaluation.final_kwh.iloc[-1]),
                         'max_violation':max(res['validation'].values())})
        comparisons=pd.DataFrame(rows);comparisons.to_csv(out/'policy_comparison.csv',index=False,encoding='utf-8-sig')
        bootstrap_comparisons(fair_daily).to_csv(out/'bootstrap_comparisons.csv',index=False,encoding='utf-8-sig')
        forecast_diagnostics(actual,forecast).to_csv(out/'forecast_diagnostics.csv',index=False,encoding='utf-8-sig')
        sensitivity=[]
        variants=[('main',100.,.9,7,1.,False),('grid50',50.,.9,7,1.,False),('scenes14',100.,.9,14,1.,False),
                  ('terminal_half',100.,.9,7,.5,False),('terminal_double',100.,.9,7,2.,False),
                  ('roundtrip90',100.,float(np.sqrt(.9)),7,1.,False),('refund',100.,.9,7,1.,True)]
        for date in SELECTED:
            d=DATES.get_loc(date);initial=float(daily.loc[daily.date==date,'initial_kwh'].iloc[0])
            for name,step,eta,count,mult,refund in variants:
                table,ds,_=simulate_day(p,actual,forecast,load_pred,d,initial,step=step,eta=eta,count=count,kappa_multiplier=mult,refund=refund)
                validate(table,eta,refund)
                sensitivity.append({'date':date,'variant':name,**{k:ds[k] for k in ['initial_kwh','final_kwh','total_cost_yuan','emergency_kwh','increase_kwh','decrease_kwh']}})
        pd.DataFrame(sensitivity).to_csv(out/'sensitivity.csv',index=False,encoding='utf-8-sig')
        summary['policy_comparison']=rows
    summary['runtime_seconds']=time.perf_counter()-begin
    (out/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':main()
