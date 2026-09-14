"""问题二：历史联合场景、日前承诺、因果储能反馈。原始附件只读。"""
from pathlib import Path
from datetime import datetime, timedelta
import argparse, hashlib, json, time
import numpy as np
import pandas as pd
import openpyxl
from scipy.optimize import linprog
from scipy.sparse import lil_matrix

ROOT = Path(__file__).resolve().parent
INPUT = Path('C:/Users/wingkii/Desktop/CUMCM2026Problems/C题/附件')
DATES = pd.date_range('2025-01-01', '2025-12-31')
SELECTED = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']
EMIN, EMAX, EREF, POWER = 1200., 10800., 6000., 5000 / 6
T = 144
LOOKBACK_WEEKS = 8  # 预测回看的同一星期几样本数；稳健性检验见模型与算法说明

def label(slot):
    m = slot * 10
    return f'{m//60:02d}:{m%60:02d}'

def interval(slot):
    return label(slot) + '-' + label(slot+1)

def minute(x):
    if hasattr(x, 'hour'):
        return x.hour*60+x.minute
    s = str(x); h, m = map(int, s.split('+')[0].split(':')[:2])
    return 60*h+m+(1440 if '+1' in s else 0)

def read_data(folder):
    w = openpyxl.load_workbook(folder/'附件1.xlsx', read_only=True, data_only=True)
    rows = list(w.active.values); w.close()
    assert len(rows) == 145 and rows[0][1] == '电价'
    assert [minute(r[0]) for r in rows[1:]] == list(range(10,1441,10))
    p = np.array([r[1] for r in rows[1:]],float)
    w = openpyxl.load_workbook(folder/'附件2.xlsx', read_only=True, data_only=True)
    arrays = []
    for name in ['小区负载','光伏发电实际功率']:
        rows = list(w[name].values)
        assert len(rows)==366 and all(len(r)==145 for r in rows)
        assert [minute(x) for x in rows[0][1:]] == list(range(10,1441,10))
        assert pd.DatetimeIndex([r[0] for r in rows[1:]]).equals(DATES)
        a = np.asarray([r[1:] for r in rows[1:]],float)/6
        assert a.shape==(365,144) and np.isfinite(a).all() and (a>=0).all()
        arrays.append(a)
    w.close()
    assert np.isfinite(p).all() and (p>0).all()
    return p, np.stack(arrays,axis=1)  # 日、负载/光伏、10分钟；单位kWh

def predict(history, date, weeks=LOOKBACK_WEEKS):
    """只接收截至昨天的历史数组，不假造调休信息。

    按精确星期几匹配：周六与周日的负载模式相反（实测19—20时相差约268 kWh），
    混同平均会得到两边都不像的曲线——周日被低估则少承诺购电、被迫按5倍价格
    买应急电，周六被高估则弃电。取最近 weeks 个同一星期几的样本平均；
    历史不足 weeks 周时用全部可得样本。
    """
    n = len(history)
    assert n > 0
    ids = np.arange(n - min(n, weeks * 7), n)
    same = ids[DATES[ids].dayofweek == date.dayofweek]
    if len(same) == 0: same = ids
    return history[same].mean(axis=0)

def scenarios(pred, residuals, history, count=7):
    if len(residuals) >= 2:
        return np.maximum(0, pred[None,:,:]+np.asarray(residuals[-count:]))
    return history[-count:].copy()

def plan_day(p, scene, initial, eta=.9, kappa=None, hard_terminal=False):
    """共享电池参考轨迹的场景LP。内部以MWh缩放，目标单位千元。"""
    if kappa is None: kappa=p[0]/eta
    s,n = scene.shape[0],len(p)
    g,c,b,e = [np.arange(k*n,(k+1)*n) for k in range(4)]
    q = np.arange(4*n,(4+s)*n).reshape(s,n)
    h=(4+s)*n; nv=h+1
    obj=np.zeros(nv);obj[g]=p;obj[q]=5*p[None,:]/s;obj[h]=kappa
    obj[c]=obj[b]=1e-8  # 小于1分钱的吞吐量择优扰动，另报告原始经济目标。
    bounds=[(0,None)]*nv
    for t in range(n):
        bounds[c[t]]=bounds[b[t]]=(0,POWER/1000)
        bounds[e[t]]=(EMIN/1000,EMAX/1000)
    if hard_terminal:
        bounds[e[-1]]=(EREF/1000,EREF/1000)
        # 独立公平边界实验：最后7段保留足够普通电源，确保实际闭环可达。
        for t in range(n-7,n): bounds[g[t]]=(POWER/1000,None)
    eq=lil_matrix((n,nv));rhs=np.zeros(n)
    for t in range(n):
        eq[t,e[t]]=1;eq[t,c[t]]=-eta;eq[t,b[t]]=1/eta
        if t: eq[t,e[t-1]]=-1
        else: rhs[t]=initial/1000
    a=lil_matrix((s*n+n+1,nv));ub=np.zeros(s*n+n+1)
    for j in range(s):
        for t in range(n):
            row=j*n+t
            a[row,g[t]]=-1;a[row,b[t]]=-1;a[row,c[t]]=1;a[row,q[j,t]]=-1
            ub[row]=(scene[j,1,t]-scene[j,0,t])/1000
    for t in range(n):
        a[s*n+t,c[t]]=1;a[s*n+t,g[t]]=-1
        ub[s*n+t]=scene[:,1,t].min()/1000
    a[-1,e[-1]]=-1;a[-1,h]=-1;ub[-1]=-EREF/1000
    res=linprog(obj,A_ub=a.tocsr(),b_ub=ub,A_eq=eq.tocsr(),b_eq=rhs,bounds=bounds,
                method='highs',options={'primal_feasibility_tolerance':1e-8,'dual_feasibility_tolerance':1e-8})
    if not res.success: raise RuntimeError(res.message)
    x=res.x*1000
    # 任意同时充放解可消去无效循环：状态不变，剩余电量增加，费用不升。
    cancel=np.minimum(x[c],x[b]/eta**2)
    charge=np.maximum(0,x[c]-cancel); discharge=np.maximum(0,x[b]-eta**2*cancel)
    grid=np.maximum(0,x[g]); storage=x[e]
    short=np.maximum(scene[:,0,:]+charge-grid-scene[:,1,:]-discharge,0)
    economic=float(p@grid+np.mean(short@(5*p))+kappa*max(EREF-storage[-1],0))
    err=max(float(np.abs(np.r_[initial,storage[:-1]]+eta*charge-discharge/eta-storage).max()),
            float(np.maximum(charge-grid-scene[:,1,:].min(axis=0),0).max()))
    assert err < 1e-4 and np.max(np.minimum(charge,discharge))<1e-5
    return grid, storage, {'plan_objective_yuan':economic,'lp_objective_yuan':float(res.fun*1000),
                           'lp_iterations':int(res.nit),'plan_violation_kwh':err}

def value_table(p, grid, scene, eta=.9, kappa=None, step=100., hard_terminal=False):
    """共享动作的边际场景价值近似；不冒充保留全部时序相关性的精确多阶段DP。"""
    if kappa is None:kappa=p[0]/eta
    states=np.arange(EMIN,EMAX+step/2,step)
    assert abs(states[-1]-EMAX)<1e-8
    delta=states[None,:]-states[:,None]
    charge=np.maximum(delta,0)/eta; discharge=np.maximum(-delta,0)*eta
    bus=charge-discharge
    valid=(charge<=POWER+1e-8)&(discharge<=POWER+1e-8)
    value=np.empty((len(p)+1,len(states)))
    value[-1]=kappa*np.maximum(EREF-states,0)
    if hard_terminal:value[-1]=np.where(np.abs(states-EREF)<1e-8,0,np.inf)
    for t in range(len(p)-1,-1,-1):
        net=scene[:,0,t]-scene[:,1,t]-grid[t]
        cost=np.maximum(net[:,None,None]+bus[None,:,:],0).mean(axis=0)*(5*p[t])
        # 仅约束普通电源可供充电；应急用于负载，约束与日前模型一致。
        allowed=valid & (charge<=grid[t]+scene[:,1,t].min()+1e-8)
        total=cost+value[t+1][None,:]
        total[~allowed]=np.inf
        value[t]=np.min(total,axis=1)
    return states,value

def action(p, grid, load, pv, energy, states, future_value, eta=.9):
    """只使用本段观测，不接收当天未来实际值；储能在网格间作分段线性插值。"""
    lo=max(EMIN,energy-POWER/eta)
    hi=min(EMAX,energy+eta*min(POWER,grid+pv))
    net=load-pv-grid
    balanced=energy-net/eta if net>=0 else energy-net*eta
    candidates=np.unique(np.r_[states[(states>=lo)&(states<=hi)],lo,hi,
                               np.clip(energy,lo,hi),np.clip(balanced,lo,hi)])
    change=candidates-energy
    c=np.maximum(change,0)/eta;b=np.maximum(-change,0)*eta
    emergency=np.maximum(net+c-b,0)
    values=5*p*emergency+np.interp(candidates,states,future_value)
    finite=np.isfinite(values)
    if not finite.any():raise RuntimeError('No feasible real-time action')
    best=values[finite].min()
    # 数值等价时最小吞吐量；不把实际吞吐惩罚计入费用。
    ids=np.flatnonzero(finite & (values<=best+1e-7))
    j=ids[np.argmin((c+b)[ids])]
    unused=max(grid+pv+b[j]+emergency[j]-load-c[j],0)
    return float(candidates[j]),float(c[j]),float(b[j]),float(emergency[j]),float(unused)

def execute(p, grid, scene, actual, initial, eta=.9, kappa=None, step=100., hard_terminal=False):
    states,value=value_table(p,grid,scene,eta,kappa,step,hard_terminal)
    energy=initial;rows=[]
    for t in range(len(p)):
        start=energy
        energy,c,b,q,w=action(p[t],grid[t],actual[0,t],actual[1,t],start,states,value[t+1],eta)
        rows.append((start,energy,c,b,q,w))
    return np.array(rows)

def baseline_grid(scene):
    # min pG+5p E[(L-V-G)+] 的经验0.8分位数；不含电池。
    x=np.sort(scene[:,0,:]-scene[:,1,:],axis=0)
    return np.maximum(0,x[int(np.ceil(.8*len(x)))-1])

def run_period(p, data, days=365, lookback_weeks=LOOKBACK_WEEKS, step=100.,
               scene_count=7, eta=.9, start_index=0, initial=EREF,
               hard_last=False, progress=False, return_scenes=False):
    energy=initial;residuals=[];records=[];daily=[];saved={}
    kappa=p[0]/eta
    # 重建真实因果残差；预测只用截至昨天的历史。
    for d in range(days):
        date=DATES[d]
        if d==0:
            pred=np.zeros((2,T));scene=None
        else:
            pred=predict(data[:d],date,lookback_weeks)
            scene=scenarios(pred,residuals,data[:d],scene_count)
        if d>=start_index:
            start=energy
            if d==0:
                grid=np.zeros(T);ref=np.full(T,energy)
                out=np.column_stack([np.full(T,energy),np.full(T,energy),np.zeros((T,2)),
                                     np.maximum(data[d,0]-data[d,1],0),np.maximum(data[d,1]-data[d,0],0)])
                info={'plan_objective_yuan':0.,'lp_objective_yuan':0.,'lp_iterations':0,'plan_violation_kwh':0.}
                basegrid=np.zeros(T)
            else:
                terminal=hard_last and d==days-1
                grid,ref,info=plan_day(p,scene,energy,eta,kappa,terminal)
                out=execute(p,grid,scene,data[d],energy,eta,kappa,step,terminal)
                basegrid=baseline_grid(scene)
            energy=out[-1,1]
            baseq=np.maximum(data[d,0]-data[d,1]-basegrid,0)
            table=pd.DataFrame(out,columns=['storage_start_kwh','storage_end_kwh','charge_kwh','discharge_kwh','emergency_kwh','unused_kwh'])
            table.insert(0,'date',date.strftime('%Y-%m-%d'));table.insert(1,'slot',np.arange(1,T+1))
            table['interval']=[interval(t) for t in range(T)]
            table['price_yuan_per_kwh']=p;table['load_kwh']=data[d,0];table['pv_kwh']=data[d,1]
            table['forecast_load_kwh']=pred[0];table['forecast_pv_kwh']=pred[1]
            table['grid_kwh']=grid;table['reference_storage_kwh']=ref
            table['plan_cost_yuan']=p*grid;table['emergency_cost_yuan']=5*p*table.emergency_kwh
            table['baseline_grid_kwh']=basegrid;table['baseline_emergency_kwh']=baseq
            table['baseline_cost_yuan']=p*basegrid+5*p*baseq
            table['lookback_weeks']=lookback_weeks;table['scenarios']=0 if scene is None else len(scene)
            table['history_latest_date']='' if d==0 else DATES[d-1].strftime('%Y-%m-%d')
            records.append(table)
            summary={'date':date.strftime('%Y-%m-%d'),'lookback_weeks':lookback_weeks,'scenarios':0 if scene is None else len(scene),
                     'initial_kwh':start,'final_kwh':energy,**info}
            for col in ['grid_kwh','emergency_kwh','charge_kwh','discharge_kwh','unused_kwh','plan_cost_yuan','emergency_cost_yuan','baseline_cost_yuan']:
                summary[col]=float(table[col].sum())
            summary['total_cost_yuan']=summary['plan_cost_yuan']+summary['emergency_cost_yuan']
            summary['prediction_mae_kwh']=float(np.mean(np.abs(pred-data[d])))
            summary['emergency_slots']=int((table.emergency_kwh>1e-6).sum())
            daily.append(summary)
            if return_scenes and str(date.date()) in SELECTED:
                saved[str(date.date())]=(pred.copy(),[r.copy() for r in residuals],start,d)
            if progress and (d%20==0 or d==days-1):
                print(f'{date.date()} completed; cost={summary["total_cost_yuan"]:.2f}; E={energy:.2f}',flush=True)
        if d>0:residuals.append(data[d]-pred)
    return pd.concat(records,ignore_index=True),pd.DataFrame(daily),saved

def validate(table, eta=.9):
    balance=table.grid_kwh+table.emergency_kwh+table.pv_kwh+table.discharge_kwh-table.load_kwh-table.charge_kwh-table.unused_kwh
    dynamic=table.storage_start_kwh+eta*table.charge_kwh-table.discharge_kwh/eta-table.storage_end_kwh
    measures={'balance_residual_kwh':float(abs(balance).max()),'storage_residual_kwh':float(abs(dynamic).max()),
              'continuity_residual_kwh':float(np.max(np.abs(table.storage_start_kwh.to_numpy()[1:]-table.storage_end_kwh.to_numpy()[:-1]))),
              'simultaneous_kwh':float(np.minimum(table.charge_kwh,table.discharge_kwh).max()),
              'charge_source_violation_kwh':float(np.maximum(table.charge_kwh-table.grid_kwh-table.pv_kwh,0).max()),
              'power_violation_kw':float(max(0,6*table[['charge_kwh','discharge_kwh']].to_numpy().max()-5000)),
              'min_storage_kwh':float(table[['storage_start_kwh','storage_end_kwh']].to_numpy().min()),
              'max_storage_kwh':float(table[['storage_start_kwh','storage_end_kwh']].to_numpy().max())}
    assert all(measures[k]<1e-4 for k in list(measures)[:6]),measures
    assert measures['min_storage_kwh']>=EMIN-1e-5 and measures['max_storage_kwh']<=EMAX+1e-5
    assert np.isfinite(table.select_dtypes('number')).all().all()
    assert (table[['grid_kwh','emergency_kwh','charge_kwh','discharge_kwh','unused_kwh']]>=-1e-6).all().all()
    return measures

def episodes(table):
    rows=[]
    for date,d in table.groupby('date',sort=False):
        q=d.emergency_kwh.to_numpy();start=None
        for t in range(T+1):
            active=t<T and q[t]>1e-6
            if active and start is None:start=t
            if not active and start is not None:
                rows.append({'date':date,'interval':label(start)+'-'+label(t),'emergency_kwh':float(q[start:t].sum())})
                start=None
    return pd.DataFrame(rows,columns=['date','interval','emergency_kwh'])

def write_excel(folder,out,table,daily,events):
    wb=openpyxl.load_workbook(folder/'附件5/result2.xlsx')
    ws=wb['计划购电量']
    for t in range(T):ws.cell(1,t+2,interval(t))
    ws.cell(1,146,'全日计划购电量（kWh）');ws.cell(1,147,'全日计划购电费（元）')
    for i,(date,d) in enumerate(table.groupby('date',sort=False),2):
        ws.cell(i,1,datetime.fromisoformat(date))
        for t,g in enumerate(d.grid_kwh,2):ws.cell(i,t,float(g))
        ws.cell(i,146,float(d.grid_kwh.sum()));ws.cell(i,147,float(d.plan_cost_yuan.sum()))
    ws.freeze_panes='B2'
    ws=wb['充放电量']
    for merged in list(ws.merged_cells.ranges):ws.unmerge_cells(str(merged))
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.value=None
    row=2
    for date,d in table.groupby('date',sort=False):
        for b in range(6):
            a=d.iloc[24*b:24*(b+1)]
            ws.cell(row,1,datetime.fromisoformat(date));ws.cell(row,2,f'{b*4:02d}:00-{(b+1)*4:02d}:00')
            ws.cell(row,3,float(a.charge_kwh.sum()));ws.cell(row,4,float(a.discharge_kwh.sum()))
            if b==0:ws.cell(row,5,'00:00');ws.cell(row,6,float(d.storage_start_kwh.iloc[0]))
            if b==1:ws.cell(row,5,'24:00');ws.cell(row,6,float(d.storage_end_kwh.iloc[-1]))
            row+=1
    ws.freeze_panes='C2'
    ws=wb['紧急购电量']
    for merged in list(ws.merged_cells.ranges):ws.unmerge_cells(str(merged))
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.value=None
    for i,r in enumerate(events.itertuples(),2):
        ws.cell(i,1,datetime.fromisoformat(r.date));ws.cell(i,2,r.interval);ws.cell(i,3,r.emergency_kwh)
    ws.freeze_panes='B2'
    ws=wb.create_sheet('每日费用汇总');ws.append(list(daily.columns))
    for row in daily.itertuples(index=False,name=None):ws.append(row)
    notes=wb.create_sheet('口径说明')
    for row in [('项目','说明'),('时间','输入按前一10分钟均值；模板标签修正为00:00至24:00。'),
                ('计划与费用','计划表仅记录0时承诺及其费用；总费用=计划费用+5倍电价紧急费用，另见每日费用汇总。'),
                ('充放电量','实际反馈执行量；非场景日前参考轨迹。每日期初等于前日期末。'),
                ('数据权限','只使用附件1电价和过去附件2；未使用附件3或未来真实曲线做日前决策。'),
                ('应急响应','假设可在本10分钟段测量负载/光伏并即时响应；未模拟段内延迟。'),
                ('近似','日前共享轨迹场景LP+100kWh储能价值网格反馈；不是一般多阶段随机控制的全局最优。'),
                ('应急时间','同日相邻正应急段合并；零应急日无记录，阈值1e-6 kWh。'),
                ('效率','充、放电效率分别0.9；能量kWh，功率上限5000kW为母线侧。')]:notes.append(row)
    for ws in wb:
        ws.column_dimensions['A'].width=19
        for row in ws:
            for cell in row:
                if isinstance(cell.value,datetime):cell.number_format='yyyy-mm-dd'
                elif isinstance(cell.value,(int,float)):cell.number_format='0.000000'
    for name in ['充放电量','紧急购电量']:
        wb[name].column_dimensions['B'].width=22
        for c in ['C','D','E','F']:wb[name].column_dimensions[c].width=20
    notes.column_dimensions['B'].width=110
    wb.save(out/'result2.xlsx')

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--input',type=Path,default=INPUT)
    parser.add_argument('--output',type=Path,default=ROOT);parser.add_argument('--days',type=int,default=365)
    parser.add_argument('--full',action='store_true',help='含公平端点对照、指定日敏感性')
    args=parser.parse_args();start_time=time.perf_counter();out=args.output.resolve()
    assert 2<=args.days<=365
    assert out!=args.input.resolve() and args.input.resolve() not in out.parents
    out.mkdir(parents=True,exist_ok=True);(out/'results').mkdir(exist_ok=True)
    p,data=read_data(args.input)
    table,daily,saved=run_period(p,data,days=args.days,lookback_weeks=LOOKBACK_WEEKS,progress=True,return_scenes=True)
    checks=validate(table)
    table.to_csv(out/'dispatch_all.csv',index=False,encoding='utf-8-sig')
    daily.to_csv(out/'daily_all.csv',index=False,encoding='utf-8-sig')
    official=table[table.date>='2025-02-01'].copy();od=daily[daily.date>='2025-02-01'].copy()
    summary={'days_run':args.days,'lookback_weeks':LOOKBACK_WEEKS,'scenario_count':7,'state_grid_kwh':100,
             'eta_charge':.9,'eta_discharge':.9,'kappa_yuan_per_internal_kwh':float(p[0]/.9),'validation':checks,
             'input_hashes':{str(f):hashlib.sha256((args.input/f).read_bytes()).hexdigest() for f in ['附件1.xlsx','附件2.xlsx','附件5/result2.xlsx']}}
    if len(official):
        events=episodes(official)
        official.to_csv(out/'dispatch.csv',index=False,encoding='utf-8-sig');od.to_csv(out/'daily.csv',index=False,encoding='utf-8-sig')
        events.to_csv(out/'emergency_events.csv',index=False,encoding='utf-8-sig')
        summary['evaluation']={c:float(od[c].sum()) for c in ['grid_kwh','emergency_kwh','charge_kwh','discharge_kwh','unused_kwh','plan_cost_yuan','emergency_cost_yuan','total_cost_yuan','baseline_cost_yuan']}
        summary['evaluation'].update({'days':len(od),'initial_kwh':float(od.initial_kwh.iloc[0]),'final_kwh':float(od.final_kwh.iloc[-1]),
                                      'emergency_slots':int(od.emergency_slots.sum()),'emergency_events':len(events)})
        summary['january_cost_yuan']=float(daily.iloc[:31].total_cost_yuan.sum())
        write_excel(args.input,out,official,od,events)
    if args.full:
        assert args.days==365
        # 同一期初6000、同一期末6000的独立对照，不改写主方案1月形成的库存。
        fair,fd,_=run_period(p,data,start_index=31,initial=EREF,hard_last=True,progress=True)
        summary['fair_boundary_validation']=validate(fair)
        assert abs(fair.storage_end_kwh.iloc[-1]-EREF)<1e-4
        fd.to_csv(out/'fair_boundary_daily.csv',index=False,encoding='utf-8-sig')
        summary['fair_boundary']={'initial_kwh':EREF,'final_kwh':float(fair.storage_end_kwh.iloc[-1]),
                                 'total_cost_yuan':float(fd.total_cost_yuan.sum()),'baseline_cost_yuan':float(fd.baseline_cost_yuan.sum())}
        sensitivity=[]
        variants=[('main',7,100.,1.,.9),('grid_50',7,50.,1.,.9),('scenes_14',14,100.,1.,.9),
                  ('scenes_28',28,100.,1.,.9),('terminal_half',7,100.,.5,.9),('terminal_double',7,100.,2.,.9),
                  ('roundtrip_90',7,100.,1.,float(np.sqrt(.9)))]
        for date,(pred,residuals,initial,d) in saved.items():
            for name,count,step,mult,eta in variants:
                scene=scenarios(pred,residuals,data[:d],count);kappa=p[0]/eta*mult
                grid,_,_=plan_day(p,scene,initial,eta,kappa)
                real=execute(p,grid,scene,data[d],initial,eta,kappa,step)
                sensitivity.append({'date':date,'variant':name,'cost_yuan':float(p@grid+5*p@real[:,4]),
                                    'emergency_kwh':float(real[:,4].sum()),'initial_kwh':initial,'final_kwh':float(real[-1,1])})
        pd.DataFrame(sensitivity).to_csv(out/'sensitivity.csv',index=False,encoding='utf-8-sig')
    summary['runtime_seconds']=time.perf_counter()-start_time
    (out/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)

if __name__=='__main__':main()
