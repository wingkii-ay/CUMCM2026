"""独立复算账本及模板输出；无外部服务、无修改原附件。"""
from pathlib import Path
import json
import numpy as np
import pandas as pd
import openpyxl

root=Path(__file__).resolve().parent
d=pd.read_csv(root/'dispatch.csv');a=pd.read_csv(root/'dispatch_all.csv')
daily=pd.read_csv(root/'daily.csv');events=pd.read_csv(root/'emergency_events.csv')
s=json.loads((root/'summary.json').read_text(encoding='utf-8'))
assert len(d)==334*144 and len(a)==365*144 and len(daily)==334
assert d.date.iloc[0]=='2025-02-01' and d.date.iloc[-1]=='2025-12-31'
assert np.all(d.groupby('date').slot.apply(lambda x:np.array_equal(x,np.arange(1,145))))
assert (pd.to_datetime(d.history_latest_date)<pd.to_datetime(d.date)).all()
assert np.max(np.abs(a.storage_start_kwh.to_numpy()[1:]-a.storage_end_kwh.to_numpy()[:-1]))<1e-6
assert abs(a.storage_start_kwh.iloc[0]-6000)<1e-6
assert np.max(abs(d.grid_kwh+d.pv_kwh+d.discharge_kwh+d.emergency_kwh-d.load_kwh-d.charge_kwh-d.unused_kwh))<1e-6
assert np.max(abs(d.storage_start_kwh+.9*d.charge_kwh-d.discharge_kwh/.9-d.storage_end_kwh))<1e-6
assert d[['storage_start_kwh','storage_end_kwh']].min().min()>=1200-1e-6
assert d[['storage_start_kwh','storage_end_kwh']].max().max()<=10800+1e-6
assert d[['charge_kwh','discharge_kwh']].max().max()*6<=5000+1e-6
assert np.minimum(d.charge_kwh,d.discharge_kwh).max()<1e-6
assert (d.charge_kwh<=d.grid_kwh+d.pv_kwh+1e-6).all()
assert np.max(abs(d.plan_cost_yuan-d.price_yuan_per_kwh*d.grid_kwh))<1e-6
assert np.max(abs(d.emergency_cost_yuan-5*d.price_yuan_per_kwh*d.emergency_kwh))<1e-6
for c in ['grid_kwh','emergency_kwh','plan_cost_yuan','emergency_cost_yuan']:
    assert abs(d[c].sum()-s['evaluation'][c])<.0001
    assert np.max(abs(d.groupby('date')[c].sum().to_numpy()-daily[c]))<1e-6
assert abs(events.emergency_kwh.sum()-d.emergency_kwh.sum())<1e-5

wb=openpyxl.load_workbook(root/'result2.xlsx',data_only=False)
ws=wb['计划购电量'];assert ws.max_row==335 and ws.max_column==147
assert ws['B1'].value=='00:00-00:10' and ws.cell(1,145).value=='23:50-24:00'
for row,(_,day) in zip(range(2,336),d.groupby('date',sort=False)):
    assert str(ws.cell(row,1).value.date())==day.date.iloc[0]
    assert np.max(abs(np.array([ws.cell(row,c).value for c in range(2,146)])-day.grid_kwh))<1e-6
    assert abs(ws.cell(row,146).value-day.grid_kwh.sum())<1e-6
    assert abs(ws.cell(row,147).value-day.plan_cost_yuan.sum())<1e-6
ws=wb['充放电量'];assert ws.max_row==334*6+1
for i,(date,day) in enumerate(d.groupby('date',sort=False)):
    for k in range(6):
        row=2+i*6+k;part=day.iloc[k*24:(k+1)*24]
        assert str(ws.cell(row,1).value.date())==date
        assert abs(ws.cell(row,3).value-part.charge_kwh.sum())<1e-6
        assert abs(ws.cell(row,4).value-part.discharge_kwh.sum())<1e-6
    assert abs(ws.cell(2+i*6,6).value-day.storage_start_kwh.iloc[0])<1e-6
    assert abs(ws.cell(3+i*6,6).value-day.storage_end_kwh.iloc[-1])<1e-6
ws=wb['紧急购电量'];assert ws.max_row==len(events)+1
for i,r in enumerate(events.itertuples(),2):
    assert str(ws.cell(i,1).value.date())==r.date and ws.cell(i,2).value==r.interval
    assert abs(ws.cell(i,3).value-r.emergency_kwh)<1e-6
    lo,hi=r.interval.split('-')
    idx=lambda x:(int(x[:2])*60+int(x[3:]))//10
    day=d[d.date==r.date];part=day.iloc[idx(lo):idx(hi)]
    assert (part.emergency_kwh>1e-6).all()
    assert abs(part.emergency_kwh.sum()-r.emergency_kwh)<1e-6
for ws in wb:
    for row in ws:
        for c in row:
            assert c.data_type!='f', 'Workbook uses values; unexpected formula requires recalculation'
            assert str(c.value) not in ['#VALUE!','#DIV/0!','#REF!','#NAME?','#NUM!','#N/A','#NULL!']
wb.close()
fair=pd.read_csv(root/'fair_boundary_daily.csv')
assert abs(fair.initial_kwh.iloc[0]-6000)<1e-6 and abs(fair.final_kwh.iloc[-1]-6000)<1e-6
assert abs(fair.total_cost_yuan.sum()-s['fair_boundary']['total_cost_yuan'])<1e-6
print('PASS: 48096 evaluation intervals; 365-day continuity; costs, physics, Excel and emergency events independently reconciled')
