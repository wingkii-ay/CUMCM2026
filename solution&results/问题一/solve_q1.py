"""问题一：确定性储能调度。python solve_q1.py [--core-only]"""
from pathlib import Path
import argparse, json, hashlib
import numpy as np
import pandas as pd
from scipy.optimize import milp, Bounds, LinearConstraint
from scipy.sparse import lil_matrix, vstack, csr_matrix
import openpyxl

INPUT = Path('D:/Users/NGALAM/OneDrive/桌面/CUMCM2026Problems/C题/附件')
ROOT = Path(__file__).resolve().parent

def read_data(folder):
    w=openpyxl.load_workbook(folder/'附件1.xlsx',read_only=True,data_only=True)
    rows=list(w.active.values);w.close()
    assert len(rows)==145 and rows[0]==('时间','电价','小区负载','光伏发电预测功率')
    for t,r in enumerate(rows[1:],1):
        x=r[0]
        if hasattr(x,'hour'): minute=x.hour*60+x.minute
        else:
            s=str(x);h,m=map(int,s.split('+')[0].split(':')[:2]);minute=h*60+m+(1440 if '+1' in s else 0)
        assert minute==10*t,(t,x)
    a=np.array([r[1:] for r in rows[1:]],float)
    assert np.isfinite(a).all() and (a>=0).all() and (a[:,0]>0).all()
    return a[:,0],a[:,1]/6,a[:,2]/6

def optimize(p,l,v,eta=.9):
    n=len(p);nv=6*n
    # 顺序：普通购电G、母线充电c、母线放电b、内部储能E、剩余w、充电模式z。
    g,c,b,e,w,z=[np.arange(k*n,(k+1)*n) for k in range(6)]
    obj=np.zeros(nv);obj[g]=p
    low=np.zeros(nv);high=np.full(nv,np.inf)
    low[e]=1200;high[e]=10800;high[c]=high[b]=5000/6;high[z]=1
    low[e[-1]]=high[e[-1]]=6000
    integrality=np.zeros(nv);integrality[z]=1
    mat=lil_matrix((5*n,nv));lb=np.full(5*n,-np.inf);ub=np.full(5*n,np.inf)
    for t in range(n):
        mat[t,g[t]]=1;mat[t,b[t]]=1;mat[t,c[t]]=-1;mat[t,w[t]]=-1
        lb[t]=ub[t]=l[t]-v[t]
        mat[n+t,e[t]]=1;mat[n+t,c[t]]=-eta;mat[n+t,b[t]]=1/eta
        if t:mat[n+t,e[t-1]]=-1;lb[n+t]=ub[n+t]=0
        else:lb[n+t]=ub[n+t]=6000
        mat[2*n+t,c[t]]=1;mat[2*n+t,z[t]]=-5000/6;ub[2*n+t]=0
        mat[3*n+t,b[t]]=1;mat[3*n+t,z[t]]=5000/6;ub[3*n+t]=5000/6
        mat[4*n+t,c[t]]=1;mat[4*n+t,g[t]]=-1;ub[4*n+t]=v[t]
    mat=mat.tocsr();constraint=LinearConstraint(mat,lb,ub);bounds=Bounds(low,high)
    options={'mip_rel_gap':1e-10,'time_limit':120}
    primary=milp(obj,integrality=integrality,bounds=bounds,constraints=constraint,options=options)
    if not primary.success:raise RuntimeError(primary.message)
    relaxed=milp(obj,integrality=np.zeros(nv),bounds=bounds,constraints=constraint)
    if not relaxed.success:raise RuntimeError(relaxed.message)
    # 多解时以母线吞吐量作为第二目标，最优费用容差为1e-7元。
    secondary=np.zeros(nv);secondary[c]=1;secondary[b]=1
    mat2=vstack([mat,csr_matrix(obj.reshape(1,-1))],format='csr')
    res=milp(secondary,integrality=integrality,bounds=bounds,
             constraints=LinearConstraint(mat2,np.r_[lb,-np.inf],np.r_[ub,primary.fun+1e-7]),options=options)
    if not res.success:raise RuntimeError(res.message)
    x=res.x;mx=mat@x
    violation=max(float(np.maximum(lb-mx,0).max()),float(np.maximum(mx-ub,0).max()),
                  float(np.maximum(low-x,0).max()),float(np.maximum(x-high,0).max()))
    assert violation<1e-5,violation
    assert np.max(np.minimum(x[c],x[b]))<1e-5
    assert abs(x[e][-1]-6000)<1e-5
    label=lambda m:f'{m//60:02d}:{m%60:02d}'
    table=pd.DataFrame({'slot':np.arange(1,n+1),'interval':[label(10*t)+'-'+label(10*t+10) for t in range(n)],
                        'price_yuan_per_kwh':p,'load_kwh':l,'pv_kwh':v,
                        'grid_kwh':x[g],'charge_kwh':x[c],'discharge_kwh':x[b],
                        'storage_start_kwh':np.r_[6000,x[e][:-1]],'storage_end_kwh':x[e],
                        'unused_kwh':x[w],'cost_yuan':p*x[g]})
    base=float(p@np.maximum(l-v,0));cost=float(obj@x)
    info={'eta_charge':eta,'eta_discharge':eta,'initial_kwh':6000,'final_kwh':float(x[e][-1]),
          'grid_kwh':float(x[g].sum()),'cost_yuan':cost,'primary_optimum_yuan':float(primary.fun),
          'milp_dual_bound_yuan':float(primary.mip_dual_bound),'mip_gap':float(primary.mip_gap),
          'lp_lower_bound_yuan':float(relaxed.fun),'no_storage_cost_yuan':base,
          'saving_yuan':base-cost,'saving_percent':100*(base-cost)/base,
          'max_constraint_violation':violation,'min_storage_kwh':float(x[e].min()),'max_storage_kwh':float(x[e].max()),
          'charge_kwh':float(x[c].sum()),'discharge_kwh':float(x[b].sum()),'unused_kwh':float(x[w].sum())}
    return table,info

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--input',type=Path,default=INPUT)
    parser.add_argument('--output',type=Path,default=ROOT);parser.add_argument('--core-only',action='store_true')
    args=parser.parse_args();out=args.output.resolve()
    assert args.input.resolve()!=out and args.input.resolve() not in out.parents
    out.mkdir(parents=True,exist_ok=True)
    p,l,v=read_data(args.input);table,info=optimize(p,l,v)
    table.to_csv(out/'dispatch.csv',index=False,encoding='utf-8-sig')
    groups=table.groupby((table.slot-1)//24)[['charge_kwh','discharge_kwh']].sum()
    groups.insert(0,'interval',[f'{4*i:02d}:00-{4*i+4:02d}:00' for i in range(6)])
    groups.to_csv(out/'storage_4h.csv',index=False,encoding='utf-8-sig')
    info['input_sha256']=hashlib.sha256((args.input/'附件1.xlsx').read_bytes()).hexdigest()
    if not args.core_only:
        _,alt=optimize(p,l,v,eta=np.sqrt(.9));info['round_trip_90_percent_sensitivity']=alt
        wb=openpyxl.load_workbook(args.input/'附件5'/'result1.xlsx')
        ws=wb['计划购电量']
        for i,row in table.iterrows():ws.cell(i+2,1,row['interval']);ws.cell(i+2,2,float(row.grid_kwh))
        ws['A147']='全天购电量（kWh）';ws['B147']=info['grid_kwh']
        ws['A148']='全天购电费（元）';ws['B148']=info['cost_yuan']
        notes=wb.create_sheet('口径说明')
        notes.append(['项目','说明'])
        notes.append(['时间映射','原模板区间整体晚10分钟；本副本按0:00–24:00和原数据右端点假设修正标签，原文件不变。'])
        notes.append(['效率','充、放电效率分别0.9；母线侧充放电量，电池内部储电量。'])
        notes.append(['功率','5000 kW按母线侧功率解释。'])
        ws=wb['充放电量']
        for i,row in groups.iterrows():
            ws.cell(int(i)+2,2,float(row.charge_kwh));ws.cell(int(i)+2,3,float(row.discharge_kwh))
        ws['D2']='0:00';ws['E2']=6000;ws['D3']='24:00';ws['E3']=info['final_kwh']
        for s in wb:
            s.column_dimensions['A'].width=28;s.column_dimensions['B'].width=20
            s.column_dimensions['C'].width=20;s.column_dimensions['D'].width=18;s.column_dimensions['E'].width=20
            for row in s:
                for cell in row:
                    if isinstance(cell.value,(float,int)):cell.number_format='0.000000'
        notes.column_dimensions['B'].width=100
        wb.save(out/'result1.xlsx')
    (out/'summary.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(info,ensure_ascii=False,indent=2))
    print(table.loc[table.slot.isin([61,73,85,97,109,121]),['interval','grid_kwh']].to_string(index=False))
    print(groups.to_string(index=False))

if __name__=='__main__':main()
