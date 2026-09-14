"""Reconcile raw inputs, constraints, workbooks, schedule versions and costs."""
import argparse,json,hashlib
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl
from core import INPUT, read_data, validate, episodes
ROOT=Path(__file__).resolve().parent

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--input',type=Path,default=INPUT);parser.add_argument('--output',type=Path,default=ROOT);args=parser.parse_args()
    out=args.output;res=out/'results';data,fc=read_data(args.input);report={};maxerr=0
    def eq(a,b):
        nonlocal maxerr
        error=float(np.max(np.abs(np.asarray(a,float)-np.asarray(b,float))));assert error<1e-5,error;maxerr=max(maxerr,error)
    for branch in [2,3]:
        df=pd.read_csv(res/f'q4{branch}_dispatch_all.csv');ev=pd.read_csv(res/f'q4{branch}_dispatch.csv')
        assert len(df)==52560 and len(ev)==48096 and ev.date.nunique()==334
        assert df.date.iloc[0]=='2025-01-01' and df.date.iloc[-1]=='2025-12-31'
        assert ev.date.iloc[0]=='2025-02-01' and ev.date.iloc[-1]=='2025-12-31'
        assert all(np.array_equal(g.slot,np.arange(1,145)) for _,g in df.groupby('date'))
        eq(df.load_kwh,data[:,0].ravel());eq(df.pv_kwh,data[:,1].ravel());eq(df.price_yuan_per_kwh,data[:,2].ravel())
        checks=validate(df);eq(df.storage_start_kwh.iloc[0],6000)
        assert (pd.to_datetime(ev.history_latest_date)<pd.to_datetime(ev.date)).all()
        assert (df.forecast_release_hour<=((df.slot-1)//36)*6).all()
        versions=pd.read_csv(res/f'q4{branch}_versions.csv');versions=versions[versions.date>='2025-02-01']
        assert (versions.slot>versions.decision_hour*6).all()
        original=versions[versions.decision_hour==0].sort_values(['date','slot'])
        latest=versions.sort_values('decision_hour').drop_duplicates(['date','slot'],keep='last').sort_values(['date','slot'])
        eq(original.version_grid_kwh,ev.original_grid_kwh);eq(latest.version_grid_kwh,ev.adjusted_grid_kwh)
        if branch==2:eq(ev.original_grid_kwh,ev.adjusted_grid_kwh)
        daily=pd.read_csv(res/f'q4{branch}_all_daily.csv')
        for col in ['total_cost_yuan','plan_cost_yuan','increase_cost_yuan','emergency_cost_yuan','emergency_kwh','original_grid_kwh','adjusted_grid_kwh']:
            eq(daily[col],ev.groupby('date')[col].sum())
        events=pd.read_csv(res/f'q4{branch}_emergency.csv');rebuilt=episodes(ev)
        assert events[['date','interval']].equals(rebuilt[['date','interval']]);eq(events.emergency_kwh,rebuilt.emergency_kwh)
        wb=openpyxl.load_workbook(out/f'result4-{branch}.xlsx',data_only=False)
        names=[('计划购电量','original_grid_kwh','plan_cost_yuan')]+([('调整购电量','adjusted_grid_kwh','increase_cost_yuan')] if branch==3 else [])
        for name,col,cost in names:
            ws=wb[name];assert ws.max_row==335 and ws.max_column==147
            assert ws.cell(1,2).value=='00:00-00:10' and ws.cell(1,145).value=='23:50-24:00'
            for i,(date,g) in enumerate(ev.groupby('date',sort=False),2):
                assert str(ws.cell(i,1).value.date())==date
                eq([ws.cell(i,j).value for j in range(2,146)],g[col]);eq(ws.cell(i,146).value,g[col].sum());eq(ws.cell(i,147).value,g[cost].sum())
        ws=wb['充放电量'];assert ws.max_row==2005
        for i,(date,g) in enumerate(ev.groupby('date',sort=False)):
            for k in range(6):
                part=g.iloc[k*24:(k+1)*24];row=2+i*6+k
                eq(ws.cell(row,3).value,part.charge_kwh.sum());eq(ws.cell(row,4).value,part.discharge_kwh.sum())
            eq(ws.cell(2+i*6,6).value,g.storage_start_kwh.iloc[0]);eq(ws.cell(3+i*6,6).value,g.storage_end_kwh.iloc[-1])
        ws=wb['紧急购电量'];assert ws.max_row==len(events)+1
        for i,row in enumerate(events.itertuples(),2):
            assert str(ws.cell(i,1).value.date())==row.date and ws.cell(i,2).value==row.interval;eq(ws.cell(i,3).value,row.emergency_kwh)
        for ws in wb:
            for row in ws:
                for cell in row:
                    assert cell.data_type not in ('f','e'),'Unexpected formula/error needs recalculation'
        wb.close();report[f'q4{branch}']={'rows':len(ev),'versions':len(versions),'events':len(events),'cost':float(ev.total_cost_yuan.sum()),'checks':checks}
    for p in res.glob('*_fair_daily.csv'):
        d=pd.read_csv(p);eq(d.initial_kwh.iloc[0],6000);eq(d.final_kwh.iloc[-1],6000);eq(d.initial_kwh.iloc[1:],d.final_kwh.iloc[:-1])
    report['max_reconciliation_error']=maxerr;report['status']='PASS'
    (res/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
