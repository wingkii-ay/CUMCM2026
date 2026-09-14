"""Run main Q4 schedules and controlled comparisons. Inputs remain read-only."""
from pathlib import Path
from datetime import datetime
from concurrent.futures import ProcessPoolExecutor
import argparse,json,time
import numpy as np
import pandas as pd
import openpyxl
from core import *
ROOT=Path(__file__).resolve().parent

def workbook(folder,out,branch,df,daily,events):
    wb=openpyxl.load_workbook(folder/'附件5'/f'result4-{branch}.xlsx')
    names=[('计划购电量','original_grid_kwh','plan_cost_yuan')]
    if branch==3:names.append(('调整购电量','adjusted_grid_kwh','increase_cost_yuan'))
    for name,col,cost in names:
        ws=wb[name]
        for t in range(144):ws.cell(1,t+2,interval(t))
        ws.cell(1,146,'全天普通电量（kWh）');ws.cell(1,147,'原计划实际费用（元）' if name=='计划购电量' else '调整相关费用（元）')
        for i,(date,g) in enumerate(df.groupby('date',sort=False),2):
            ws.cell(i,1,datetime.fromisoformat(date))
            for j,value in enumerate(g[col],2):ws.cell(i,j,float(value)).number_format='0.000000'
            ws.cell(i,146,float(g[col].sum())).number_format='0.000000';ws.cell(i,147,float(g[cost].sum())).number_format='0.00'
        ws.freeze_panes='B2'
    ws=wb['充放电量']
    for m in list(ws.merged_cells.ranges):ws.unmerge_cells(str(m))
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.value=None
    for i,(date,g) in enumerate(df.groupby('date',sort=False)):
        for k in range(6):
            part=g.iloc[k*24:(k+1)*24];values=[datetime.fromisoformat(date),label(k*24)+'-'+label((k+1)*24),float(part.charge_kwh.sum()),float(part.discharge_kwh.sum())]
            if k==0:values+=['00:00',float(g.storage_start_kwh.iloc[0])]
            if k==1:values+=['24:00',float(g.storage_end_kwh.iloc[-1])]
            for j,value in enumerate(values,1):ws.cell(2+i*6+k,j,value)
    ws=wb['紧急购电量']
    for m in list(ws.merged_cells.ranges):ws.unmerge_cells(str(m))
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.value=None
    for i,r in enumerate(events.itertuples(),2):
        ws.cell(i,1,datetime.fromisoformat(r.date));ws.cell(i,2,r.interval);ws.cell(i,3,r.emergency_kwh)
    ws=wb.create_sheet('每日实际费用');ws.append(list(daily.columns))
    for row in daily.itertuples(index=False,name=None):ws.append(row)
    ws=wb.create_sheet('口径说明')
    for row in [('项目','说明'),('电价','未来未知；规划用历史预测，执行按附件4供电时段实际价结算，不锁价。'),
                ('费用','计划全额收费；调整仅最终A对原G净差计一次；应急5倍；终端罚不计账单。'),
                ('电量','母线侧kWh；计划G，调整A不含应急；余量可含光伏及购电剩余。'),
                ('时间','00:10解释为前10分钟均值，输出标签覆盖00:00至24:00。'),
                ('储能','两侧效率各0.9；功率5000kW；内部1200至10800kWh；从1月1日6000连续推进。'),
                ('预测','负载8个同星期几，电价28日均值及已观测偏差修正；7历史联合残差场景。'),
                ('算法','保守共享轨迹LP与100kWh网格实际反馈；0时对6时起可调整时段的缺额按合同增购价1.5倍定阶，之前时段按5倍应急；不是全年多阶段全局最优。'),
                ('库存','主结果不每日恢复6000；公平边界对照单列，不覆盖主结果。')]:ws.append(row)
    for ws in wb:
        ws.column_dimensions['A'].width=20
        for row in ws:
            for cell in row:
                if isinstance(cell.value,datetime):cell.number_format='yyyy-mm-dd'
                elif isinstance(cell.value,(int,float)) and cell.number_format=='General':cell.number_format='0.000000'
    wb['口径说明'].column_dimensions['B'].width=110
    wb.save(out/f'result4-{branch}.xlsx');wb.close()

def run_case(payload):
    data,fc,base,prices,branch,policy,fair=payload
    rows=[];days=[];versions=[];energy=6000.;started=time.time()
    for d in range(31 if fair else 0,365):
        df,day,ver=simulate_day(data,fc,base,prices,d,energy,branch,policy,hard=fair and d==364)
        energy=day['final_kwh'];rows.append(df);days.append(day);versions.extend(ver)
        if d%90==0:print(f'Q4-{branch} {policy} fair={fair} {DATES[d].date()}',flush=True)
    table=pd.concat(rows,ignore_index=True);checks=validate(table)
    if fair:assert abs(energy-6000)<1e-5
    return dict(branch=branch,policy=policy,fair=fair,table=table,daily=pd.DataFrame(days),versions=versions,checks=checks,seconds=time.time()-started)

def smoke(folder,out):
    out.mkdir(exist_ok=True,parents=True);data,fc=read_data(folder);base,prices=predictions(data);report={}
    for branch in [2,3]:
        for d,hard in [(31,False),(364,True)]:
            df,day,_=simulate_day(data,fc,base,prices,d,6000,branch,hard=hard)
            report[f'b{branch}_d{d}']={**day,'checks':validate(df)}
            if hard:assert abs(day['final_kwh']-6000)<1e-5
    # Modify today's afternoon and ALL future actual L,V,p. Before noon must stay unchanged.
    d=78;alt=data.copy();alt[d,:,72:]*=1.7;alt[d+1:]*=1.3;ab,ap=predictions(alt)
    for branch in [2,3]:
        a,_,_=simulate_day(data,fc,base,prices,d,6000,branch)
        b,_,_=simulate_day(alt,fc,ab,ap,d,6000,branch)
        original_error=float(abs(a.original_grid_kwh-b.original_grid_kwh).max())
        prefix_error=float(abs(a.iloc[:72][['storage_end_kwh','adjusted_grid_kwh']]-b.iloc[:72][['storage_end_kwh','adjusted_grid_kwh']]).max().max())
        assert max(original_error,prefix_error)<1e-7
        report[f'causal_b{branch}']=dict(original_error=original_error,prefix_error=prefix_error)
    # Price-demand covariance: minimum initial inventory prevents free discharge.
    sc=np.zeros((2,3,1));sc[:,0,0]=[0,100];sc[:,2,0]=[1,10]
    g,e,val=plan(sc,1200,0);assert abs(g[0]-100)<1e-5 and abs(val-550)<1e-5
    sc[:,2,0]=[10,1];g,e,val=plan(sc,1200,0)
    assert abs(g[0])<1e-5 and abs(val-250)<1e-5
    report['price_demand_covariance']='PASS: high-price demand -> G100/cost550; low-price demand -> G0/cost250'
    report['status']='PASS';(out/'smoke.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--input',type=Path,default=INPUT);parser.add_argument('--output',type=Path,default=ROOT)
    parser.add_argument('--smoke',action='store_true');parser.add_argument('--workers',type=int,default=2);args=parser.parse_args()
    if args.smoke:return smoke(args.input,args.output)
    out=args.output;out.mkdir(exist_ok=True,parents=True);res=out/'results';res.mkdir(exist_ok=True)
    data,fc=read_data(args.input);base,prices=predictions(data)
    configs=[(2,'all',False),(3,'all',False),(2,'all',True),(3,'all',True),(3,'only0',True),(3,'stale_all',True),(3,'stale18',True)]
    jobs=[(data,fc,base,prices,*c) for c in configs];summary={};fair={};mains={}
    with ProcessPoolExecutor(max_workers=args.workers) as pool:
        for result in pool.map(run_case,jobs):
            branch,policy,isfair=result['branch'],result['policy'],result['fair'];key=f'q4{branch}_{policy}'+('_fair' if isfair else '')
            table=result['table'];daily=result['daily'];daily.to_csv(res/f'{key}_daily_all.csv',index=False,encoding='utf-8-sig')
            evaluation=daily[daily.date>='2025-02-01'];evaluation.to_csv(res/f'{key}_daily.csv',index=False,encoding='utf-8-sig')
            ev=table[table.date>='2025-02-01'].reset_index(drop=True)
            summary[key]={'cost':float(evaluation.total_cost_yuan.sum()),'initial':float(evaluation.initial_kwh.iloc[0]),'final':float(evaluation.final_kwh.iloc[-1]),
                          'emergency':float(evaluation.emergency_kwh.sum()),'checks':result['checks'],'seconds':result['seconds']}
            if isfair:fair[key]=evaluation
            else:
                mains[branch]=table
                table.to_csv(res/f'q4{branch}_dispatch_all.csv',index=False,encoding='utf-8-sig');ev.to_csv(res/f'q4{branch}_dispatch.csv',index=False,encoding='utf-8-sig')
                events=episodes(ev);events.to_csv(res/f'q4{branch}_emergency.csv',index=False,encoding='utf-8-sig')
                versions=pd.DataFrame(result['versions'],columns=['date','decision_hour','forecast_release_hour','slot','original_grid_kwh','version_grid_kwh'])
                versions.to_csv(res/f'q4{branch}_versions.csv',index=False,encoding='utf-8-sig')
                workbook(args.input,out,branch,ev,evaluation,events)
            (res/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
            print(key,summary[key],flush=True)
    rng=np.random.default_rng(2026);comparisons=[]
    for before,after in [('q43_only0_fair','q43_all_fair'),('q43_stale_all_fair','q43_all_fair'),('q43_stale18_fair','q43_all_fair'),('q42_all_fair','q43_all_fair')]:
        x=fair[before].total_cost_yuan.to_numpy()-fair[after].total_cost_yuan.to_numpy();n=len(x)
        starts=rng.integers(0,n,size=(2000,int(np.ceil(n/7))));idx=((starts[:,:,None]+np.arange(7))%n).reshape(2000,-1)[:,:n]
        lo,hi=np.quantile(x[idx].mean(axis=1),[.025,.975])
        comparisons.append(dict(before=before,after=after,saving=float(x.sum()),mean=float(x.mean()),ci_low=lo,ci_high=hi))
    pd.DataFrame(comparisons).to_csv(res/'comparisons.csv',index=False,encoding='utf-8-sig')
    sensitivity=[]
    for branch,table in mains.items():
        for date in SELECTED:
            d=int((pd.Timestamp(date)-DATES[0]).days);initial=float(table[table.date==date].storage_start_kwh.iloc[0])
            for name,step,eta in [('main',100,.9),('grid50',50,.9),('roundtrip90',100,np.sqrt(.9))]:
                df,day,_=simulate_day(data,fc,base,prices,d,initial,branch,step=step,eta=eta);validate(df,eta)
                sensitivity.append({'case':name,**day})
    pd.DataFrame(sensitivity).to_csv(res/'sensitivity.csv',index=False,encoding='utf-8-sig')
    diagnostics=[]
    for r in range(4):
        off=r*36
        for d in range(31,365):
            err=prices[d,r,off:]-data[d,2,off:]
            diagnostics.append(dict(date=str(DATES[d].date()),release_hour=r*6,price_mae=float(abs(err).mean()),price_rmse=float(np.sqrt(np.mean(err**2)))))
    pd.DataFrame(diagnostics).to_csv(res/'price_diagnostics.csv',index=False,encoding='utf-8-sig')

if __name__=='__main__':main()
